"""Excel workbook export of an assessment + summary."""
import io

from .util import to_gib


def _j(items, f):
    return "; ".join(f(x) for x in items or [])


def _disk(d):
    return f"{d.get('name')} {to_gib(d.get('size_bytes'))}GiB {d.get('model') or ''}".strip()


def _nic(n):
    s = f"{n.get('name')} {n.get('mac') or ''}"
    if n.get("speed_mbps"):
        s += f" {n['speed_mbps'] / 1000:g}Gb"
    if n.get("ipv4"):
        s += f" {n['ipv4']}"
    return s


def tables(assessment, summary):
    t = {}
    tot = summary["totals"]
    t["Summary"] = (["Metric", "Value"], [
        ["Generated", summary["generated_at"]],
        ["OpenShift clusters", len(assessment.get("openshift", []))],
        ["OpenShift core-pair subscriptions (virtual workers)", tot["ocp_core_pair_subs"]],
        ["OpenShift socket-pair subscriptions (bare-metal workers)", tot["ocp_bm_socket_pair_subs"]],
        ["OpenShift workers unclassified (physical/virtual)", tot["ocp_unknown_nodes"]],
        ["ODF core-pair subscriptions (storage nodes)", tot["odf_core_pair_subs"]],
        ["RHEL physical hosts", summary["rhel"]["physical"]["hosts"]],
        ["RHEL physical socket-pair subscriptions", tot["rhel_physical_subs"]],
        ["RHEL VMs on VMware (powered on)", summary["rhel"]["vmware"]["vms"]],
        ["RHEL VMware - per-guest subscriptions (2 VMs each)", tot["rhel_vmware_per_guest_subs"]],
        ["RHEL VMware - alternative: Virtual Datacenter subscriptions", tot["rhel_vmware_vdc_subs"]],
        ["RHV hypervisor hosts", summary["rhel"]["rhv"]["hypervisor_hosts"]],
        ["RHV hypervisor socket-pair subscriptions", tot["rhv_hypervisor_socket_pair_subs"]],
        ["RHEL VMs on RHV (powered on)", summary["rhel"]["rhv"]["vms"]],
        ["RHEL RHV - per-guest subscriptions (2 VMs each)", tot["rhel_rhv_per_guest_subs"]],
        ["RHEL RHV - alternative: Virtual Datacenter subscriptions", tot["rhel_rhv_vdc_subs"]],
        ["RHEL VMs on other hypervisors", summary["rhel"]["other_virtual"]["vms"]],
        ["RHEL other-virtual subscriptions", tot["rhel_other_virtual_subs"]],
        ["RHEL VMs on OpenShift Virtualization", tot["ocpv_rhel_vms"]],
        [], ["Notes", ""], *[[n, ""] for n in summary["notes"]],
    ])
    rows = []
    for c in summary["openshift"]:
        ch = c["chargeable"]
        rows.append([c["name"], c["version"], c["platform"], c["api_url"], sum(r["nodes"] for r in c["roles"].values()),
                     ch["nodes"], ch["virtual_vcpu"], ch["core_pair_subs"], ch["bm_nodes"], ch["bm_sockets"],
                     ch["bm_socket_pair_subs"], ch["unknown_nodes"], c["odf"]["installed"], c["odf"]["core_pair_subs"],
                     c["virtualization"]["installed"], c["virtualization"]["vms"], c["virtualization"]["rhel_vms"]])
    t["OCP Clusters"] = (["Cluster", "Version", "Platform", "API", "Nodes", "Chargeable nodes", "Virtual worker vCPU",
                          "Core-pair subs", "BM worker nodes", "BM sockets", "BM socket-pair subs", "Unclassified workers",
                          "ODF", "ODF core-pair subs", "OCP Virt", "VMs", "RHEL VMs"], rows)
    rows = []
    for c in assessment.get("openshift", []):
        for n in c.get("nodes", []):
            hw = n.get("hw") or {}
            rows.append([c.get("name"), n["name"], n["role"], ",".join(n["roles"]), n["ready"], n["virtualization"], n["internal_ip"],
                         n["os_image"], n["arch"], n["cpu_capacity"], hw.get("sockets"), hw.get("physical_cores"),
                         hw.get("cpu_model"), to_gib(n["memory_bytes"]), to_gib(n.get("disk_total_bytes")),
                         len(hw.get("disks") or []), _j(hw.get("disks"), _disk), n.get("nic_count"), _j(hw.get("nics"), _nic),
                         hw.get("vendor"), hw.get("model"), hw.get("serial"), ",".join(hw.get("sources") or [])])
    t["OCP Nodes"] = (["Cluster", "Node", "Role", "Labels", "Ready", "Virtualization", "IP", "OS", "Arch", "vCPU/Logical CPU",
                       "Sockets", "Physical cores", "CPU model", "Memory GiB", "Disk total GiB", "Disk count", "Disks",
                       "NIC count", "NICs", "Vendor", "Model", "Serial", "HW source"], rows)
    rows = []
    for c in assessment.get("openshift", []):
        for v in (c.get("virtualization") or {}).get("vms", []):
            rows.append([c.get("name"), v["namespace"], v["name"], v["status"], v["node"], v["os_family"], v["os_version"],
                         v["os_name"], v["os_source"], v.get("vcpu"), to_gib(v.get("memory_bytes")), to_gib(v.get("disk_total_bytes")),
                         _j(v.get("disks"), lambda d: f"{d.get('name')}:{to_gib(d.get('size_bytes'))}GiB"),
                         _j(v.get("nics"), lambda n: f"{n.get('name')}({n.get('network')}) {n.get('ipv4') or ''}"), v.get("instancetype")])
    t["OCP Virt VMs"] = (["Cluster", "Namespace", "VM", "Status", "Node", "OS family", "OS ver", "OS", "OS source", "vCPU",
                          "Memory GiB", "Disk GiB", "Disks", "NICs", "Instancetype"], rows)
    rows = []
    for c in assessment.get("openshift", []):
        for o in c.get("operators", []):
            rows.append([c.get("name"), o["display_name"], o["name"], o["version"], o["namespace"], o["provider"], o["phase"]])
    t["OCP Operators"] = (["Cluster", "Operator", "CSV", "Version", "Namespace", "Provider", "Phase"], rows)
    rows = []
    for vc in assessment.get("vcenter", []):
        for h in vc.get("hosts", []):
            rows.append([vc.get("name"), h["datacenter"], h["cluster"], h["name"], h["esxi_version"], h["vendor"], h["model"],
                         h["cpu_model"], h["sockets"], h["cores"], h["threads"], to_gib(h["memory_bytes"]), len(h["nics"]),
                         _j(h["nics"], _nic), h["vm_count"], h["connection_state"], h["maintenance"]])
    t["vCenter Hosts"] = (["vCenter", "Datacenter", "Cluster", "Host", "ESXi", "Vendor", "Model", "CPU model", "Sockets", "Cores",
                           "Threads", "Memory GiB", "NIC count", "NICs", "VMs", "Connection", "Maintenance"], rows)
    rows = []
    for vc in assessment.get("vcenter", []):
        for v in vc.get("vms", []):
            rows.append([vc.get("name"), v["datacenter"], v["cluster"], v["esxi_host"], v["name"], v["power_state"], v["os_family"],
                         v["os_version"], v["os_name"], v["os_source"], v["config_mismatch"], v["hostname"], v["ip"], v["vcpu"],
                         v["cores_per_socket"], to_gib(v["memory_bytes"]), len(v["disks"]), to_gib(v["disk_total_bytes"]),
                         _j(v["disks"], lambda d: f"{d['name']} {to_gib(d['size_bytes'])}GiB [{d['datastore']}]"),
                         len(v["nics"]), _j(v["nics"], lambda n: f"{n['model']} {n['network']} {n.get('ipv4') or ''}"),
                         v["tools_status"]])
    t["vCenter VMs"] = (["vCenter", "Datacenter", "Cluster", "ESXi host", "VM", "Power", "OS family", "OS ver", "Guest OS",
                         "OS source", "Config/Tools OS mismatch", "Hostname", "IP", "vCPU", "Cores/socket", "Memory GiB",
                         "Disk count", "Disk GiB", "Disks", "NIC count", "NICs", "Tools"], rows)
    rows = []
    for rv in assessment.get("rhv", []) or []:
        for h in rv.get("hosts", []):
            rows.append([rv.get("name"), h["datacenter"], h["cluster"], h["name"], h["address"], h["type"], h["os"], h["status"],
                         h["vendor"], h["model"], h["serial"], h["cpu_model"], h["sockets"], h["cores"], h["threads"],
                         to_gib(h["memory_bytes"]), len(h["nics"]), _j(h["nics"], _nic), h["vms_active"], h["hosted_engine"]])
    t["RHV Hosts"] = (["RHV Manager", "Data center", "Cluster", "Host", "Address", "Type", "OS", "Status", "Vendor", "Model",
                       "Serial", "CPU model", "Sockets", "Cores", "Threads", "Memory GiB", "NIC count", "NICs", "Running VMs",
                       "Hosted-engine host"], rows)
    rows = []
    for rv in assessment.get("rhv", []) or []:
        for v in rv.get("vms", []):
            rows.append([rv.get("name"), v["datacenter"], v["cluster"], v["host"], v["name"], v["status"], v["os_family"],
                         v["os_version"], v["os_name"], v["os_source"], v["os_type_config"], v["hosted_engine"], v["hostname"], v["ip"],
                         v["vcpu"], f"{v['sockets']}x{v['cores_per_socket']}x{v['threads_per_core']}", to_gib(v["memory_bytes"]),
                         len(v["disks"]), to_gib(v["disk_total_bytes"]), to_gib(v["disk_actual_bytes"]),
                         _j(v["disks"], lambda d: f"{d['name']} {to_gib(d['size_bytes'])}GiB [{d['storage_domain']}] {d['interface']}"),
                         len(v["nics"]), _j(v["nics"], lambda n: f"{n['name']} {n['model']} {n['network']} {n['mac']} {n.get('ipv4') or ''}"),
                         v["vm_type"], v["description"]])
    t["RHV VMs"] = (["RHV Manager", "Data center", "Cluster", "Host", "VM", "Status", "OS family", "OS ver", "Guest OS", "OS source",
                     "Configured OS type", "HostedEngine", "FQDN", "IP", "vCPU", "Topology (SxCxT)", "Memory GiB", "Disk count",
                     "Disk provisioned GiB", "Disk actual GiB", "Disks", "NIC count", "NICs", "VM type", "Description"], rows)
    rows = []
    for b in assessment.get("rhel", []):
        for h in b.get("hosts", []):
            cpu, sm = h.get("cpu") or {}, h.get("subscription") or {}
            rows.append([h["target"], h.get("status"), h.get("error", ""), h.get("hostname"), h.get("os_name"), h.get("kernel"),
                         h.get("arch"), h.get("virtualization"), (h.get("hardware") or {}).get("vendor"), (h.get("hardware") or {}).get("model"),
                         (h.get("hardware") or {}).get("serial"), cpu.get("model"), cpu.get("sockets"), cpu.get("cores_per_socket"),
                         cpu.get("physical_cores"), cpu.get("logical_cpus"), to_gib(h.get("memory_bytes")), len(h.get("disks") or []),
                         to_gib(h.get("disk_total_bytes")), _j(h.get("disks"), _disk), len(h.get("nics") or []), _j(h.get("nics"), _nic),
                         sm.get("overall_status"), sm.get("sca"), sm.get("registered_to"), sm.get("org"), ",".join(h.get("addons") or [])])
    t["RHEL Hosts (SSH)"] = (["Target", "Status", "Error", "Hostname", "OS", "Kernel", "Arch", "Virtualization", "Vendor", "Model",
                              "Serial", "CPU model", "Sockets", "Cores/socket", "Physical cores", "Logical CPUs", "Memory GiB",
                              "Disk count", "Disk GiB", "Disks", "NIC count", "NICs", "Subscription status", "SCA",
                              "Registered to", "Org", "Add-ons"], rows)
    rows = [[e["name"], ",".join(e["sources"]), e.get("os_version"), "Virtual" if e["virtual"] else "Physical", e.get("platform"),
             e.get("cluster"), e.get("power_state"), e.get("sockets"), e.get("vcpu"), to_gib(e.get("memory_bytes")),
             e.get("subscription"), e.get("registered_to")] for e in summary["rhel"]["estate"]]
    t["RHEL Estate"] = (["Name", "Sources", "RHEL ver", "Type", "Platform", "Cluster", "Power", "Sockets", "vCPU", "Memory GiB",
                         "Subscription status", "Registered to"], rows)
    t["RHEL VDC by Cluster"] = (["vCenter", "Datacenter", "Cluster", "Hosts", "Sockets", "RHEL VMs", "Per-guest subs", "VDC subs"],
                                [[x["vcenter"], x["datacenter"], x["cluster"], x["hosts"], x["sockets"], x["rhel_vms"],
                                  x["per_guest_subs"], x["vdc_subs"]] for x in summary["rhel"]["vmware"]["vdc_clusters"]])
    t["RHEL RHV VDC by Cluster"] = (["RHV Manager", "Data center", "Cluster", "Hosts", "Sockets", "RHEL VMs", "Per-guest subs", "VDC subs"],
                                    [[x["rhv"], x["datacenter"], x["cluster"], x["hosts"], x["sockets"], x["rhel_vms"],
                                      x["per_guest_subs"], x["vdc_subs"]] for x in summary["rhel"]["rhv"]["vdc_clusters"]])
    return t


def to_xlsx(assessment, summary):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    wb.remove(wb.active)
    hdr_fill = PatternFill("solid", fgColor="A30000")
    for name, (cols, rows) in tables(assessment, summary).items():
        ws = wb.create_sheet(name[:31])
        ws.append(cols)
        for r in rows:
            ws.append([("" if v is None else v) for v in r])
        for c in ws[1]:
            c.font = Font(bold=True, color="FFFFFF")
            c.fill = hdr_fill
            c.alignment = Alignment(vertical="center", wrap_text=True)
        ws.freeze_panes = "A2"
        if rows and name != "Summary":
            ws.auto_filter.ref = ws.dimensions
        for i, col in enumerate(cols, 1):
            width = max([len(str(col))] + [len(str(r[i - 1])) for r in rows[:300] if len(r) >= i and r[i - 1] is not None])
            ws.column_dimensions[get_column_letter(i)].width = min(max(10, width + 2), 70 if name == "Summary" else 50)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
